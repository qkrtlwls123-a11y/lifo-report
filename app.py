# -*- coding: utf-8 -*-
"""
LIFO 진단 결과 → PPT 자동 생성 웹 앱
기존 템플릿(CIAM 진단_LIFO_Report.pptx)의 서식을 그대로 복제하여 사용
"""
import os
import re
import uuid
import time
import tempfile
from copy import deepcopy
from datetime import datetime
from flask import Flask, render_template, request, jsonify, send_file
import openpyxl
from pypdf import PdfReader
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from lxml import etree

# LIFO 유형 순서 (S/G, C/T, C/H, A/D)
LIFO_KEYS = ['SG', 'CT', 'CH', 'AD']

app = Flask(__name__)

# 배포 환경: /tmp 사용, 로컬: uploads/ 사용
if os.environ.get('RENDER') or os.environ.get('RAILWAY_ENVIRONMENT'):
    UPLOAD_DIR = os.path.join(tempfile.gettempdir(), 'lifo-uploads')
else:
    UPLOAD_DIR = os.path.join(os.path.dirname(__file__), 'uploads')

app.config['UPLOAD_FOLDER'] = UPLOAD_DIR
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024  # 50MB
TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), 'template.pptx')

os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)


def cleanup_old_files():
    """30분 이상 지난 생성 파일을 자동 삭제합니다."""
    now = time.time()
    for f in os.listdir(app.config['UPLOAD_FOLDER']):
        fp = os.path.join(app.config['UPLOAD_FOLDER'], f)
        if os.path.isfile(fp) and now - os.path.getmtime(fp) > 1800:
            os.remove(fp)


# ── 엑셀 파싱 ───────────────────────────────────────────────
def parse_lifo_excel(file_path, original_filename=''):
    """엑셀 파일에서 LIFO 진단 점수를 추출합니다."""
    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheets = wb.sheetnames

    if len(sheets) < 2:
        raise ValueError(f"시트가 2개 이상이어야 합니다. (현재: {len(sheets)}개)")

    ws = wb[sheets[1]]  # 두 번째 시트 = 결과표

    # TOTAL+ (Row 12): D12=S+, F12=C+, H12=C+, J12=A+
    s_plus = ws['D12'].value or 0
    c_plus = ws['F12'].value or 0
    ch_plus = ws['H12'].value or 0
    a_plus = ws['J12'].value or 0

    # TOTAL- (Row 13): D13=G-, F13=T-, H13=H-, J13=D-
    g_minus = ws['D13'].value or 0
    t_minus = ws['F13'].value or 0
    h_minus = ws['H13'].value or 0
    d_minus = ws['J13'].value or 0

    # 파일명에서 이름/부서 추출
    name, dept = extract_name_from_filename(original_filename)

    # 순위 결과 (I5) - 수식 오류일 수 있으므로 직접 계산
    top_type = ws['I5'].value or ''
    if not top_type or str(top_type).startswith('#') or str(top_type) == 'None':
        type_map = {'SG': s_plus, 'CT': c_plus, 'CH': ch_plus, 'AD': a_plus}
        top_key = max(type_map, key=type_map.get)
        top_type = top_key

    wb.close()

    return {
        'name': name,
        'dept': dept,
        'top_type': str(top_type).strip(),
        'plus': {
            'SG': int(s_plus), 'CT': int(c_plus),
            'CH': int(ch_plus), 'AD': int(a_plus),
            'total': int(s_plus) + int(c_plus) + int(ch_plus) + int(a_plus)
        },
        'minus': {
            'SG': int(g_minus), 'CT': int(t_minus),
            'CH': int(h_minus), 'AD': int(d_minus),
            'total': int(g_minus) + int(t_minus) + int(h_minus) + int(d_minus)
        },
    }


# ── PDF 파싱 ────────────────────────────────────────────────
def parse_lifo_pdf(file_path, original_filename=''):
    """PDF 진단 결과지에서 LIFO 점수를 추출합니다.

    엑셀과 동일한 데이터를 담은 PDF 리포트를 지원합니다.
    - "TOTAL +" / "TOTAL -" 행에서 4개 유형 점수를 읽습니다.
    - "나의 유형"에서 1순위 유형을 읽습니다.
    """
    reader = PdfReader(file_path)
    text = "\n".join((page.extract_text() or "") for page in reader.pages)

    plus = _parse_total_line(text, '+')
    minus = _parse_total_line(text, '-')

    # TOTAL 행을 못 찾으면 세부 항목(A~L / a~l) 합산으로 대체
    if plus is None:
        plus = _parse_component_scores(text, upper=True)
    if minus is None:
        minus = _parse_component_scores(text, upper=False)

    if plus is None or minus is None:
        raise ValueError('PDF에서 LIFO 점수를 찾을 수 없습니다. (TOTAL +/- 행 확인)')

    name, dept = extract_name_from_filename(original_filename)

    # 1순위 유형: "나의 유형" 근처에서 추출, 없으면 최고 점수로 계산
    m = re.search(r'(SG|CT|CH|AD)\s*나의\s*유형|나의\s*유형\s*(SG|CT|CH|AD)', text)
    top_type = (m.group(1) or m.group(2)) if m else max(plus, key=plus.get)

    return {
        'name': name,
        'dept': dept,
        'top_type': str(top_type).strip(),
        'plus': {
            **{k: int(plus[k]) for k in LIFO_KEYS},
            'total': sum(int(plus[k]) for k in LIFO_KEYS),
        },
        'minus': {
            **{k: int(minus[k]) for k in LIFO_KEYS},
            'total': sum(int(minus[k]) for k in LIFO_KEYS),
        },
    }


def _parse_total_line(text, sign):
    """'TOTAL +' 또는 'TOTAL -' 행에서 4개 유형 점수를 추출합니다."""
    marker = 'TOTAL +' if sign == '+' else 'TOTAL -'
    for line in text.splitlines():
        if marker in line:
            pairs = re.findall(r'([A-Za-z])\s+(\d+)', line)
            nums = [int(n) for _, n in pairs]
            if len(nums) >= 4:
                return dict(zip(LIFO_KEYS, nums[:4]))
    return None


def _parse_component_scores(text, upper=True):
    """세부 항목 점수(A~L 또는 a~l)를 합산하여 4개 유형 점수를 계산합니다.

    LIFO 채점 규칙:
      SG = A+E+I,  CT = B+F+J,  CH = C+G+K,  AD = D+H+L  (긍정, 대문자)
      SG = a+e+i,  CT = b+f+j,  CH = c+g+k,  AD = d+h+l  (부정, 소문자)
    """
    letters = 'ABCDEFGHIJKL' if upper else 'abcdefghijkl'
    scores = {}
    for line in text.splitlines():
        for letter, val in re.findall(r'(?:^|\s)([A-La-l])\s+(\d+)(?:\s|$)', line):
            if letter in letters and letter not in scores:
                scores[letter] = int(val)
    if len(scores) < 12:
        return None
    groups = {
        'SG': letters[0::4][:3],   # A/E/I
        'CT': letters[1::4][:3],   # B/F/J
        'CH': letters[2::4][:3],   # C/G/K
        'AD': letters[3::4][:3],   # D/H/L
    }
    return {k: sum(scores[c] for c in cols) for k, cols in groups.items()}


# 파일명에서 제거할 라벨 토큰 (LIFO 리포트 관련 단어 / 날짜)
FILENAME_LABEL_RE = re.compile(
    r'^(LIFO|진단지|진단|결과지|결과|리포트|레포트|보고서|report|\d{4,})',
    re.IGNORECASE,
)


def extract_name_from_filename(filename):
    """파일명에서 이름과 부서를 추출합니다.

    지원 패턴:
      - ...진단지_부서_이름              (엑셀 진단지)
      - 이름_LIFO진단, 부서_이름_LIFO진단  (PDF 리포트)
    """
    base = os.path.splitext(filename)[0]

    # 패턴: ...진단지_부서_이름 (이름에 공백 포함 가능, 예: "이 웅", "정청산 책임")
    match = re.search(r'진단지[_\s]+(.+?)_([^_]+?)\s*$', base)
    if match:
        return match.group(2).strip(), match.group(1).strip()

    # 패턴2: 진단지_부서 이름 (언더스코어 없이 공백으로 구분)
    match2 = re.search(r'진단지[_\s]+(.+?)\s+(\S+)\s*$', base)
    if match2:
        return match2.group(2).strip(), match2.group(1).strip()

    # 'LIFO진단', '결과', 날짜 등 라벨 필드를 제거하고 남는 것을 부서/이름으로 사용
    # 예: '홍길동_LIFO진단' → 이름 '홍길동', '영업부_홍길동_LIFO진단' → 부서 '영업부', 이름 '홍길동'
    # (필드 구분은 '_'만 사용해 '이 웅'처럼 공백이 들어간 이름을 보존)
    tokens = [t.strip() for t in base.split('_') if t.strip()]
    kept = [t for t in tokens if not FILENAME_LABEL_RE.search(t)]
    if len(kept) >= 2:
        return kept[-1].strip(), kept[-2].strip()
    if len(kept) == 1:
        return kept[0].strip(), ''

    # 폴백: 언더스코어 분리
    parts = re.split(r'[_]', base)
    if len(parts) >= 2:
        return parts[-1].strip(), parts[-2].strip() if len(parts) >= 3 else ''
    return base, ''


# ── PPT 생성 (템플릿 복제 방식) ─────────────────────────────
def clone_slide(prs, template_slide):
    """템플릿 슬라이드를 복제합니다. 모든 도형/서식/스타일이 그대로 유지됩니다."""
    slide_layout = template_slide.slide_layout
    new_slide = prs.slides.add_slide(slide_layout)

    # 새 슬라이드의 기본 도형 제거
    spTree = new_slide.shapes._spTree
    for sp in list(spTree):
        tag = sp.tag.split('}')[-1] if '}' in sp.tag else sp.tag
        if tag in ('sp', 'cxnSp', 'graphicFrame', 'pic', 'grpSp'):
            spTree.remove(sp)

    # 템플릿의 모든 도형을 깊은 복사
    for child in template_slide.shapes._spTree:
        tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
        if tag in ('sp', 'cxnSp', 'graphicFrame', 'pic', 'grpSp'):
            spTree.append(deepcopy(child))

    return new_slide


def find_tables(slide):
    """슬라이드에서 테이블 도형을 위치(top) 기준으로 정렬하여 반환합니다."""
    tables = []
    for shape in slide.shapes:
        if shape.has_table:
            tables.append(shape)
    tables.sort(key=lambda s: s.top)
    return tables


def update_cell_text_preserve_style(cell, new_text):
    """셀 텍스트를 변경하되, 기존 서식(폰트, 크기, 굵기 등)을 모두 유지합니다."""
    ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
    runs = cell._tc.findall('.//a:r', ns)
    if runs:
        # 첫 번째 run의 텍스트만 교체, 나머지 run 제거
        t_elem = runs[0].find('a:t', ns)
        if t_elem is not None:
            t_elem.text = str(new_text)
        for extra_run in runs[1:]:
            extra_run.getparent().remove(extra_run)
    else:
        # run이 없으면 새로 생성 (기본 스타일)
        cell.text = str(new_text)


def update_score_color(cell, is_red):
    """점수 셀의 글자색을 변경합니다. 빨간색=#FF0000, 일반=#000000"""
    ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
    A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
    rPr = cell._tc.find('.//a:rPr', ns)
    if rPr is not None:
        # 기존 solidFill 제거
        old_fill = rPr.find('a:solidFill', ns)
        if old_fill is not None:
            rPr.remove(old_fill)
        # 새 색상을 rPr의 첫 번째 자식으로 삽입 (OOXML 순서 준수)
        color_val = 'FF0000' if is_red else '000000'
        solidFill = etree.Element(f'{{{A}}}solidFill')
        srgbClr = etree.SubElement(solidFill, f'{{{A}}}srgbClr')
        srgbClr.set('val', color_val)
        rPr.insert(0, solidFill)


def check_red_condition(plus_score, minus_score):
    """엑셀 조건부 서식 규칙을 적용합니다.
    조건1: 점수 >= 28 → 빨간 글씨
    조건2: +와 -의 차이 > 4이고 자기가 더 큰 쪽 → 빨간 글씨
    """
    diff = abs(plus_score - minus_score)
    plus_red = plus_score >= 28 or (diff > 4 and plus_score > minus_score)
    minus_red = minus_score >= 28 or (diff > 4 and minus_score > plus_score)
    return plus_red, minus_red


def fill_table_data(table_shape, person_data):
    """테이블에 한 사람의 LIFO 데이터를 채웁니다."""
    table = table_shape.table

    # [0,0] = 이름
    update_cell_text_preserve_style(table.cell(0, 0), person_data['name'])

    plus = person_data['plus']
    minus = person_data['minus']
    keys = ['SG', 'CT', 'CH', 'AD']

    for i, key in enumerate(keys):
        p_score = plus[key]
        m_score = minus[key]
        p_red, m_red = check_red_condition(p_score, m_score)

        # + 행 [1, i+1]
        cell_plus = table.cell(1, i + 1)
        update_cell_text_preserve_style(cell_plus, str(p_score))
        update_score_color(cell_plus, p_red)

        # - 행 [2, i+1]
        cell_minus = table.cell(2, i + 1)
        update_cell_text_preserve_style(cell_minus, str(m_score))
        update_score_color(cell_minus, m_red)


def generate_pptx(participants):
    """참가자 리스트를 받아, 기존 템플릿을 복제하여 PPT를 생성합니다."""
    prs = Presentation(TEMPLATE_PATH)

    # 템플릿 슬라이드 1 (데이터가 채워진 예시)을 복제 원본으로 사용
    template_slide = prs.slides[1]

    # 필요한 슬라이드 수 계산 (2명/슬라이드)
    num_slides_needed = (len(participants) + 1) // 2

    # 슬라이드 복제 생성
    new_slides = []
    for _ in range(num_slides_needed):
        new_slide = clone_slide(prs, template_slide)
        new_slides.append(new_slide)

    # 각 슬라이드에 데이터 채우기
    for slide_idx, slide in enumerate(new_slides):
        tables = find_tables(slide)
        if len(tables) < 2:
            continue

        # 상단 테이블 = 첫 번째 사람
        p1_idx = slide_idx * 2
        if p1_idx < len(participants):
            fill_table_data(tables[0], participants[p1_idx])

        # 하단 테이블 = 두 번째 사람
        p2_idx = slide_idx * 2 + 1
        if p2_idx < len(participants):
            fill_table_data(tables[1], participants[p2_idx])
        else:
            # 홀수 명일 때 하단 테이블 비우기
            clear_table(tables[1])

    # 원본 템플릿 슬라이드 2개 삭제 (인덱스 0, 1)
    delete_slide(prs, 1)
    delete_slide(prs, 0)

    # 저장
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    filename = f'LIFO_Report_{timestamp}.pptx'
    filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
    prs.save(filepath)

    return filepath, filename


def clear_table(table_shape):
    """테이블의 데이터 셀을 비웁니다."""
    table = table_shape.table
    # 이름 비우기
    update_cell_text_preserve_style(table.cell(0, 0), '')
    # 점수 비우기
    for ri in range(1, 3):
        for ci in range(1, 5):
            update_cell_text_preserve_style(table.cell(ri, ci), '')
            # 색상도 기본 검정으로
            update_score_color(table.cell(ri, ci), False)


def delete_slide(prs, slide_index):
    """프레젠테이션에서 슬라이드를 삭제합니다."""
    rId = prs.slides._sldIdLst[slide_index].get(
        '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
    )
    prs.part.drop_rel(rId)
    del prs.slides._sldIdLst[slide_index]


# ── Flask 라우트 ─────────────────────────────────────────────
@app.route('/')
def index():
    return render_template('index.html')


@app.route('/upload', methods=['POST'])
def upload():
    """엑셀 또는 PDF 파일을 업로드하고 LIFO 점수를 파싱합니다."""
    if 'files' not in request.files:
        return jsonify({'error': '파일이 없습니다.'}), 400

    files = request.files.getlist('files')
    results = []
    errors = []

    for f in files:
        if not f.filename:
            continue
        lower = f.filename.lower()
        if lower.endswith(('.xlsx', '.xls')):
            ext, parser = '.xlsx', parse_lifo_excel
        elif lower.endswith('.pdf'):
            ext, parser = '.pdf', parse_lifo_pdf
        else:
            errors.append(f'{f.filename}: 엑셀(.xlsx) 또는 PDF 파일이 아닙니다.')
            continue

        temp_path = os.path.join(app.config['UPLOAD_FOLDER'], f'{uuid.uuid4().hex}{ext}')
        try:
            f.save(temp_path)
            data = parser(temp_path, f.filename)
            data['filename'] = f.filename
            results.append(data)
        except Exception as e:
            errors.append(f'{f.filename}: {str(e)}')
        finally:
            if os.path.exists(temp_path):
                os.remove(temp_path)

    return jsonify({'results': results, 'errors': errors})


@app.route('/generate', methods=['POST'])
def generate():
    """파싱된 데이터를 받아 PPT를 생성합니다."""
    data = request.get_json()
    if not data or 'participants' not in data:
        return jsonify({'error': '데이터가 없습니다.'}), 400

    participants = data['participants']
    if not participants:
        return jsonify({'error': '참가자가 없습니다.'}), 400

    try:
        cleanup_old_files()
        filepath, filename = generate_pptx(participants)
        return jsonify({'filename': filename, 'download_url': f'/download/{filename}'})
    except Exception as e:
        import traceback
        return jsonify({'error': str(e), 'trace': traceback.format_exc()}), 500


@app.route('/download/<filename>')
def download(filename):
    """생성된 PPT를 다운로드합니다."""
    filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
    if not os.path.exists(filepath):
        return jsonify({'error': '파일을 찾을 수 없습니다.'}), 404
    return send_file(filepath, as_attachment=True, download_name=filename)


if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    debug = not os.environ.get('RENDER')
    app.run(debug=debug, host='0.0.0.0', port=port)
